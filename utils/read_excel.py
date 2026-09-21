import openpyxl


def read_test_cases(file_path):
    """读取 Excel 测试用例，返回字典列表"""
    wb = openpyxl.load_workbook(file_path)
    ws = wb.active

    headers = [cell.value for cell in ws[1]]

    cases = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        case = dict(zip(headers, row))
        cases.append(case)

    return cases


if __name__ == "__main__":
    cases = read_test_cases("../data/test_cases.xlsx")
    for case in cases:
        print(case)